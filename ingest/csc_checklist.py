#!/usr/bin/env python3
"""Ingest the official CSC Certification Action Checklist XLSX into the actions table.

Source: NYSDEC Climate Smart Communities portal (climatesmart.ny.gov), sheet
"December 18, 2025 (version 4.5)". Structure:
  - header rows 2-4 explain tags (CEC, CSC Grants)
  - row 5 is the column header (Action Name | Legacy Action # | Legacy Action Name | Points | Type/Status)
  - pledge-element section rows look like "PE1: Build a climate-smart community."
  - action rows look like "PE1 Action: CSC Task Force" / "Innovation: ..." / "Performance: ..."

Idempotent: upserts on (program_slug, external_id). Re-runnable whenever DEC
publishes a new annual checklist.
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl
import requests

from db import CSC_BASE_URL, BROWSER_UA, connect, slugify, upsert_action, upsert_program

CHECKLIST_URL = (
    f"{CSC_BASE_URL}/fileadmin/csc/documents/Action_Checklist/2025-12-18-CSC-CertActionChecklist.xlsx"
)
ACTIONS_PAGE_URL = f"{CSC_BASE_URL}/actions-certification/actions/"
RAW_DIR = Path(__file__).resolve().parent / "data" / "raw"

SECTION_RE = re.compile(r"^PE(\d+):\s*(.+)$")
ACTION_RE = re.compile(r"^(PE(\d+)\s+Action|Innovation|Performance):\s*(.+)$")
RANGE_RE = re.compile(r"^(\d+)\s*-\s*(\d+)$")


def download(dest: Path, refresh: bool) -> None:
    if dest.exists() and not refresh:
        print(f"[checklist] using cached download: {dest}")
        return
    dest.parent.mkdir(parents=True, exist_ok=True)
    print(f"[checklist] downloading {CHECKLIST_URL}")
    resp = requests.get(CHECKLIST_URL, headers={"User-Agent": BROWSER_UA}, timeout=60)
    resp.raise_for_status()
    dest.write_bytes(resp.content)
    print(f"[checklist] saved {len(resp.content)} bytes -> {dest}")


def parse_points(raw: str | None) -> tuple[int | None, int | None, str, bool]:
    """Return (min, max, display, variable). Handles '20', '12 - 16', '2- 4'."""
    if raw is None:
        return None, None, "", False
    display = str(raw).strip()
    m = RANGE_RE.match(display)
    if m:
        lo, hi = int(m.group(1)), int(m.group(2))
        return min(lo, hi), max(lo, hi), display, True
    if display.isdigit():
        n = int(display)
        return n, n, display, False
    # unexpected format: keep display, no numeric bounds
    print(f"[checklist] WARN unrecognized points value: {display!r}")
    return None, None, display, False


def parse_rows(ws) -> list[dict]:
    # find the header row
    header_row = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        if row and str(row[0]).strip() == "Action Name":
            header_row = i
            break
    if header_row is None:
        raise SystemExit("[checklist] could not find the 'Action Name' header row")

    actions: list[dict] = []
    seen: dict[str, int] = {}
    pledge_element = ""
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        col_a = str(row[0]).strip() if row[0] is not None else ""

        # section headers
        m = SECTION_RE.match(col_a)
        if m:
            pledge_element = f"PE{m.group(1)}"
            continue
        if col_a in ("Innovation", "Performance"):
            pledge_element = col_a
            continue

        # action rows
        m = ACTION_RE.match(col_a)
        if not m:
            continue  # blank rows / footers
        name = str(m.group(3)).strip()
        legacy_no = str(row[1]).strip() if row[1] is not None else None
        legacy_name = str(row[2]).strip() if row[2] is not None else None
        legacy_no = None if legacy_no in ("", "NA", "None") else legacy_no
        legacy_name = None if legacy_name in ("", "NA", "None") else legacy_name

        points_min, points_max, points_display, variable = parse_points(row[3])
        tags = [t.strip() for t in str(row[4]).split(",") if t.strip()] if row[4] else []

        base = slugify(name)
        n = seen.get(base, 0)
        external_id = base if n == 0 else f"{base}-{n + 1}"  # deterministic on re-run
        seen[base] = n + 1

        actions.append(
            {
                "program_slug": "csc",
                "external_id": external_id,
                "name": name,
                "pledge_element": pledge_element,
                "legacy_action_no": legacy_no,
                "legacy_action_name": legacy_name,
                "points_min": points_min,
                "points_max": points_max,
                "points_display": points_display,
                "variable_points": variable,
                "mandatory": "Mandatory" in tags,
                "priority": "Priority" in tags,
                "funding_eligible": "CSC Grants" in tags,
                "cec_related": "CEC" in tags,
                "status_tags": tags,
                "source_url": ACTIONS_PAGE_URL,
            }
        )
    return actions

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--refresh", action="store_true", help="re-download the XLSX even if cached")
    args = ap.parse_args()

    dest = RAW_DIR / "2025-12-18-CSC-CertActionChecklist.xlsx"
    download(dest, args.refresh)

    wb = openpyxl.load_workbook(dest, data_only=True)
    ws = wb.active
    actions = parse_rows(ws)
    if not actions:
        raise SystemExit("[checklist] no actions parsed — sheet format changed?")

    with connect() as conn:
        upsert_program(
            conn,
            "csc",
            "Climate Smart Communities",
            "NYSDEC",
            "active",
            CSC_BASE_URL,
        )
        for row in actions:
            upsert_action(conn, row)
        conn.commit()

    by_pe = Counter(a["pledge_element"] for a in actions)
    variable = sum(1 for a in actions if a["variable_points"])
    print(f"[checklist] ingested {len(actions)} actions:")
    for pe in sorted(by_pe):
        print(f"    {pe:>11}: {by_pe[pe]}")
    print(f"    variable-point actions: {variable}")
    print(f"    mandatory: {sum(1 for a in actions if a['mandatory'])}")
    print(f"    funding-eligible (CSC Grants): {sum(1 for a in actions if a['funding_eligible'])}")
    print(f"    CEC-related: {sum(1 for a in actions if a['cec_related'])}")
    return 0


if __name__ == "__main__":
    sys.exit(main())